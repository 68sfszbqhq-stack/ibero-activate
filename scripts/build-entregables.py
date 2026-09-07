# -*- coding: utf-8 -*-
"""
Exporta el tablero a Excel (.xlsx con colores) y a PDF (la matriz completa
en una hoja ancha).

    python3 scripts/build-entregables.py

Lee scripts/datos-tablero.json, la misma fuente del tablero y del .docx.
Las calificaciones salen de `precarga`, que es exactamente lo que trae el
tablero al abrirlo. Si José ya cambió cosas en el navegador, esas viven en su
localStorage y se exportan desde el propio tablero con el botón Excel.
"""

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATOS = os.path.join(RAIZ, 'scripts', 'datos-tablero.json')
XLSX = os.path.join(RAIZ, 'Tablero de áreas - Pausas Activas.xlsx')
PDF = os.path.join(RAIZ, 'Tablero de áreas - Pausas Activas.pdf')

LETRA = {'g': 'si', 'm': 'mas', 'n': 'no'}
ETIQUETA = {'si': 'Le gusta', 'mas': 'Más o menos', 'no': 'No le gusta', '': ''}
HEX = {'si': '16A34A', 'mas': '2563EB', 'no': 'DC2626', '': 'F3F4F6'}
RGB = {k: colors.HexColor('#' + v) for k, v in HEX.items()}

ROJO = '9A1B2F'
TINTA = '101014'
GRIS = '6B7280'


def cargar():
    if not os.path.exists(DATOS):
        sys.exit('✗ Falta scripts/datos-tablero.json. Corre antes: node scripts/build-tablero.js')
    d = json.load(open(DATOS, encoding='utf-8'))

    # Matriz [actividad][area] -> 'si' | 'mas' | 'no' | ''
    valores = {}
    for a in d['actividades']:
        fila = d.get('precarga', {}).get(a['id'], '')
        valores[a['id']] = {
            ar['id']: LETRA.get(fila[i], '') if i < len(fila) else ''
            for i, ar in enumerate(d['areas'])
        }
    return d, valores


# ─────────────────────────────── EXCEL ───────────────────────────────
def excel(d, valores):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tablero'

    borde = Border(*[Side(style='thin', color='E5E7EB')] * 4)
    centro = Alignment(horizontal='center', vertical='center')
    vertical = Alignment(horizontal='center', vertical='bottom', textRotation=90)

    # Encabezado
    ws.cell(1, 1, 'Actividad').font = Font(bold=True, color=TINTA, size=11)
    ws.cell(1, 2, 'Aspecto').font = Font(bold=True, color=TINTA, size=11)
    ws.cell(1, 3, 'Segmento').font = Font(bold=True, color=TINTA, size=11)
    ws.cell(1, 4, 'Min').font = Font(bold=True, color=TINTA, size=11)
    for j, ar in enumerate(d['areas']):
        c = ws.cell(1, 5 + j, ar['nombre'])
        c.font = Font(bold=True, size=9, color=TINTA)
        c.alignment = vertical
    ws.row_dimensions[1].height = 150

    # Filas, agrupadas por segmento en el orden del documento
    fila = 2
    for sk in d['orden']:
        seg = d['segmentos'][sk]
        acts = [a for a in d['actividades'] if a['segmento'] == sk]
        if not acts:
            continue

        c = ws.cell(fila, 1, seg['nombre'])
        c.font = Font(bold=True, size=10, color=ROJO)
        for j in range(1, 5 + len(d['areas'])):
            ws.cell(fila, j).fill = PatternFill('solid', fgColor='F6F7F9')
        fila += 1

        for a in sorted(acts, key=lambda x: x['nombre'].lower()):
            ws.cell(fila, 1, a['nombre']).font = Font(size=10)
            ws.cell(fila, 2, d['aspectos'][a['aspecto']]['nombre']).font = Font(size=9, color=GRIS)
            ws.cell(fila, 3, seg['nombre']).font = Font(size=9, color=GRIS)
            ws.cell(fila, 4, a.get('duracion') or '').alignment = centro
            for j, ar in enumerate(d['areas']):
                v = valores[a['id']][ar['id']]
                cel = ws.cell(fila, 5 + j, ETIQUETA[v])
                cel.fill = PatternFill('solid', fgColor=HEX[v])
                cel.font = Font(size=8, color='FFFFFF' if v else '9CA3AF')
                cel.alignment = centro
                cel.border = borde
            fila += 1

    ws.freeze_panes = 'E2'
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 6
    for j in range(len(d['areas'])):
        ws.column_dimensions[get_column_letter(5 + j)].width = 12

    # Hoja de referencia: qué tan participativa es cada área
    ws2 = wb.create_sheet('Áreas')
    for k, t in enumerate(['Área', 'Personas', 'Asistencias', 'Por persona', 'Participación'], 1):
        ws2.cell(1, k, t).font = Font(bold=True)
    for i, ar in enumerate(sorted(d['areas'], key=lambda x: -x.get('porPersona', 0)), 2):
        ws2.cell(i, 1, ar['nombre'])
        ws2.cell(i, 2, ar.get('personas', 0))
        ws2.cell(i, 3, ar.get('asistencias', 0))
        ws2.cell(i, 4, ar.get('porPersona', 0))
        ws2.cell(i, 5, {'alta': 'Alta', 'media': 'Media', 'baja': 'Baja'}.get(ar.get('banda'), ''))
    for col, w in zip('ABCDE', [42, 10, 12, 12, 14]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = 'A2'

    wb.save(XLSX)
    return fila - 2


# ──────────────────────────────── PDF ────────────────────────────────
def pdf(d, valores):
    areas = d['areas']
    filas = []
    for sk in d['orden']:
        acts = [a for a in d['actividades'] if a['segmento'] == sk]
        if not acts:
            continue
        filas.append(('seg', d['segmentos'][sk]['nombre']))
        for a in sorted(acts, key=lambda x: x['nombre'].lower()):
            filas.append(('act', a))

    ANCHO_NOM = 46 * mm
    CELDA_W, CELDA_H = 7.2 * mm, 5.0 * mm
    MARGEN = 10 * mm
    CAB = 42 * mm     # espacio para los nombres de área, escritos de canto

    ancho = MARGEN * 2 + ANCHO_NOM + CELDA_W * len(areas)
    alto = MARGEN * 2 + CAB + CELDA_H * len(filas) + 16 * mm

    c = canvas.Canvas(PDF, pagesize=(ancho, alto))
    c.setTitle('Tablero de áreas · Pausas Activas')

    y = alto - MARGEN
    c.setFillColor(colors.HexColor('#101014'))
    c.setFont('Helvetica-Bold', 13)
    c.drawString(MARGEN, y - 4 * mm, 'Tablero de áreas · Pausas Activas')
    c.setFont('Helvetica', 8.5)
    c.setFillColor(colors.HexColor('#6B7280'))
    c.drawString(MARGEN, y - 9 * mm,
                 f"{len([f for f in filas if f[0] == 'act'])} actividades × {len(areas)} áreas · "
                 'Movimiento con Sentido Ignaciano · IBERO Puebla')

    # Leyenda
    lx = MARGEN
    ly = y - 15 * mm
    for v, t in [('si', 'Le gusta'), ('mas', 'Más o menos'), ('no', 'No le gusta')]:
        c.setFillColor(RGB[v])
        c.rect(lx, ly, 3.2 * mm, 3.2 * mm, stroke=0, fill=1)
        c.setFillColor(colors.HexColor('#101014'))
        c.setFont('Helvetica', 8)
        c.drawString(lx + 4.5 * mm, ly + 0.6 * mm, t)
        lx += 26 * mm

    tope = ly - 4 * mm

    # Nombres de área, de canto
    c.setFont('Helvetica', 6.6)
    c.setFillColor(colors.HexColor('#101014'))
    for j, ar in enumerate(areas):
        x = MARGEN + ANCHO_NOM + j * CELDA_W + CELDA_W / 2
        c.saveState()
        c.translate(x + 2, tope - CAB + 2 * mm)
        c.rotate(90)
        c.drawString(0, 0, ar['corto'][:30])
        c.restoreState()

    yy = tope - CAB
    for tipo, dato in filas:
        yy -= CELDA_H
        if tipo == 'seg':
            c.setFillColor(colors.HexColor('#FBF4F5'))
            c.rect(MARGEN, yy, ANCHO_NOM + CELDA_W * len(areas), CELDA_H, stroke=0, fill=1)
            c.setFillColor(colors.HexColor('#9A1B2F'))
            c.setFont('Helvetica-Bold', 6.8)
            c.drawString(MARGEN + 1.5 * mm, yy + 1.5 * mm, dato)
            continue

        c.setFillColor(colors.HexColor('#101014'))
        c.setFont('Helvetica', 6.8)
        nombre = dato['nombre']
        while c.stringWidth(nombre, 'Helvetica', 6.8) > ANCHO_NOM - 8 * mm and len(nombre) > 4:
            nombre = nombre[:-1]
        c.drawString(MARGEN + 3 * mm, yy + 1.5 * mm, nombre)
        if dato.get('duracion'):
            c.setFillColor(colors.HexColor('#9CA3AF'))
            c.drawRightString(MARGEN + ANCHO_NOM - 1.5 * mm, yy + 1.5 * mm, f"{dato['duracion']}'")

        for j, ar in enumerate(areas):
            v = valores[dato['id']][ar['id']]
            c.setFillColor(RGB[v])
            c.rect(MARGEN + ANCHO_NOM + j * CELDA_W + 0.5,
                   yy + 0.5, CELDA_W - 1.6, CELDA_H - 1.6, stroke=0, fill=1)

    c.showPage()
    c.save()
    return len(filas)


def main():
    d, valores = cargar()
    n = excel(d, valores)
    pdf(d, valores)

    cuenta = {'si': 0, 'mas': 0, 'no': 0, '': 0}
    for f in valores.values():
        for v in f.values():
            cuenta[v] += 1

    print(f'✓ {os.path.basename(XLSX)}')
    print(f'✓ {os.path.basename(PDF)}')
    print(f'   {len(d["actividades"])} actividades × {len(d["areas"])} áreas = '
          f'{len(d["actividades"]) * len(d["areas"])} cuadritos')
    print(f'   {cuenta["si"]} le gusta · {cuenta["mas"]} más o menos · '
          f'{cuenta["no"]} no le gusta · {cuenta[""]} vacíos')


if __name__ == '__main__':
    main()
